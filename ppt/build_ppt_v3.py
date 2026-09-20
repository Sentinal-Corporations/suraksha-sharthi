from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE = r"C:\Users\Admin\Downloads\bharat innovation 2.0\CHANAKYA_SurakshaParchi"
TEAL = RGBColor(14,124,123); SAFFRON = RGBColor(255,122,26); INK = RGBColor(11,31,34)
WHITE = RGBColor(255,255,255); LIGHT = RGBColor(244,247,246); GREY = RGBColor(220,228,227)
RED = RGBColor(215,38,61); MUT = RGBColor(110,120,120)

prs = Presentation(); prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def base(bg, notes=""):
    sl = prs.slides.add_slide(BLANK)
    fl = sl.background.fill; fl.solid(); fl.fore_color.rgb = bg
    bar = sl.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(0.09))
    bar.fill.solid(); bar.fill.fore_color.rgb = SAFFRON; bar.line.fill.background()
    ft = sl.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.35))
    p = ft.text_frame.paragraphs[0]
    p.text = "Team CHANAKYA  •  Suraksha Parchi — Clear Parchi, Safe Patient  •  ayush4ru@gmail.com"
    p.font.size = Pt(11); p.font.color.rgb = MUT
    if notes: sl.notes_slide.placeholders[1].text = notes
    return sl

def kick(sl, t, dark=False):
    tx = sl.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.5))
    p = tx.text_frame.paragraphs[0]; p.text = t; p.font.size = Pt(15); p.font.bold = True
    p.font.color.rgb = SAFFRON if dark else TEAL

def head(sl, t, dark=False):
    tx = sl.shapes.add_textbox(Inches(0.6), Inches(0.72), Inches(7.6), Inches(1.0))
    tx.text_frame.word_wrap = True
    p = tx.text_frame.paragraphs[0]; p.text = t; p.font.size = Pt(29); p.font.bold = True
    p.font.color.rgb = WHITE if dark else INK

def bullets(sl, l, t, w, h, items, size=17, color=INK):
    tx = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)); tf = tx.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + it; p.font.size = Pt(size); p.font.color.rgb = color
        p.space_after = Pt(6); p.space_before = Pt(2)

def pic(sl, name, l, t, w, h):
    try: sl.shapes.add_picture(f"{BASE}/{name}", Inches(l), Inches(t), Inches(w), Inches(h))
    except Exception as e: print("img miss", name, e)

def cap(sl, l, t, w, text, color=MUT):
    tx = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(0.7)); tx.text_frame.word_wrap = True
    p = tx.text_frame.paragraphs[0]; p.text = text; p.font.size = Pt(12); p.font.color.rgb = color

# 1 COVER
sl = base(INK, "Open with the hand-raise question. One line on who we are, then the tagline. Keep it warm, not robotic.")
pic(sl, "branding/logo.png", 8.9, 0.5, 3.7, 3.7)
tx = sl.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(7.8), Inches(0.5))
p = tx.text_frame.paragraphs[0]; p.text = "BHARAT INNOVATION CHALLENGE 2.0  •  HEALTHCARE & HEALTHTECH"
p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = SAFFRON
tx2 = sl.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(7.8), Inches(2.1)); tx2.text_frame.word_wrap = True
p2 = tx2.text_frame.paragraphs[0]; p2.text = "Suraksha Parchi\nClear Parchi, Safe Patient"; p2.font.size = Pt(46); p2.font.bold = True; p2.font.color.rgb = WHITE
bullets(sl, 0.6, 3.5, 7.8, 3.1, ["Team CHANAKYA — Ayush Pandey (Leader) • Kaushal Rawat (CMO) • +2 joining", "Mentor: [To Fill]  •  ayush4ru@gmail.com  •  Hindi • Bengali • Marathi", "Photo lo, Bhasha me samjho — a working app, not a concept"], 16, WHITE)

# 2 PROBLEM
sl = base(LIGHT, "Tell dadi's story first, then the numbers. Jury remembers stories, slides prove them.")
kick(sl, "2  PROBLEM & OPPORTUNITY"); head(sl, "Nobody can read the parchi")
bullets(sl, 0.6, 1.95, 7.4, 4.9, ["60%+ prescriptions still handwritten in English — elders and migrants just guess", "Half of all patients take the wrong dose; 1L+ deaths a year from medication errors", "Half-finished antibiotics are fuelling drug resistance across India", "10L+ ASHA workers carry village healthcare — with zero tools for this"], 18)
pic(sl, "branding/sample-parchi.png", 8.5, 1.9, 3.9, 2.8)
cap(sl, 8.5, 4.85, 3.9, "A real parchi: Crocin + Azithro + HbA1c 8.2 — to most families, just scribble.")

# 3 GAPS
sl = base(LIGHT, "Be fair to competitors, then land the punch: nobody does all five things we do.")
kick(sl, "3  EXISTING SOLUTIONS & GAP ANALYSIS"); head(sl, "Apps exist. Not for Bharat.")
bullets(sl, 0.6, 1.95, 7.4, 4.9, ["City health apps need English + internet — and ignore handwriting completely", "Basic scanners read print, fail on cursive; our test: barely half the drugs right", "Symptom chatbots guess doses and hallucinate — dangerous with medicines", "Nobody combines: messy-handwriting read + mother-tongue voice + pictograms + duplicate alerts + ASHA offline mode"], 18)
pic(sl, "mvp/shots/07_desktop.png", 8.5, 1.9, 3.9, 2.6)
cap(sl, 8.5, 4.6, 3.9, "Our answer runs as a real phone app — full-screen on mobile, sandbox on desktop.")

# 4 SOLUTION
sl = base(LIGHT, "Walk the four steps slowly. This is the slide the jury photographs.")
kick(sl, "4  PROPOSED SOLUTION"); head(sl, "Photo lo → samjho → yaad rakho")
bullets(sl, 0.6, 1.95, 12.1, 1.5, ["Scan the parchi → hybrid AI reads it → explains disease + dose by voice → timetable reminds, red flags warn"], 18)
pic(sl, "branding/journey.png", 0.6, 3.4, 12.1, 3.2)

# 5 TECH
sl = base(LIGHT, "Lead with Needle-2 numbers — downloaded and tested, not claimed. Offline-first is the differentiator.")
kick(sl, "5  TECHNOLOGY & INNOVATION"); head(sl, "14MB of AI that lives in the phone")
bullets(sl, 0.6, 1.95, 7.4, 4.9, ["Needle-2 (45M params, 13.74MB) — downloaded, wired and tested on our medicine tools", "Hybrid read: instant on-device scan, smart cloud pass only when handwriting is messy", "Grounded answers: 49-medicine verified DB + 12 interaction rules — the AI can never invent a dose", "Voice in Hindi, Bengali, Marathi (Bhashini + offline fallback); works with zero internet"], 17)
pic(sl, "branding/architecture.png", 8.5, 1.95, 3.9, 1.35)
pic(sl, "branding/logo.png", 8.5, 3.5, 3.9, 2.7)

# 6 PROTOTYPE
sl = base(LIGHT, "Demo live here: scan the sample, fire the duplicate alert. If WiFi dies, play the narrated video.")
kick(sl, "6  PROTOTYPE / LIVE DEMONSTRATION"); head(sl, "Working app — try to break it")
bullets(sl, 0.6, 1.95, 6.6, 4.9, ["Real flow: sample parchi → 87% read → Hindi voice + pictogram timetable", "Crowd moment: add Dolo beside Crocin — RED duplicate-overdose alert fires", "Suvidha tab: Jan Aushadhi prices, Ayushman packages, one-tap helplines", "Backup: 5-minute narrated Hinglish demo video (demo_video_app.mp4)"], 17)
pic(sl, "mvp/shots/03_result.png", 7.6, 1.7, 2.2, 3.9)
pic(sl, "mvp/shots/04_alert.png", 10.0, 1.7, 2.2, 3.9)

# 7 USERS
sl = base(LIGHT, "Name the humans: dadi, the migrant worker, the ASHA didi. Then the scale path.")
kick(sl, "7  TARGET USERS & USE CASES"); head(sl, "Built for dadi. Ready for 100Cr.")
bullets(sl, 0.6, 1.95, 7.4, 4.9, ["Elders and low-literate families — voice + pictures, zero English needed", "Migrants — Bengali and Marathi out of the box, more languages plug in", "10L+ ASHA workers — doorstep scans, referral slips, adherence ticks", "Jan Aushadhi counters — scan once, print the timetable in 3 languages"], 18)
pic(sl, "mvp/shots/05_suvidha.png", 8.5, 1.7, 3.4, 3.8)

# 8 IMPACT
sl = base(LIGHT, "Every claim ties to something the Neon dashboard will actually measure in the pilot.")
kick(sl, "8  IMPACT & OUTCOMES"); head(sl, "Fewer wrong doses, finished courses")
bullets(sl, 0.6, 1.95, 7.4, 4.9, ["Adherence +40% via voice + reminders — measured per family in the pilot", "Antibiotic course completion +35% — a direct hit on AMR", "Duplicate and overdose alerts catch emergencies before they happen", "Sugar/BP cross-checks catch uncontrolled cases weeks earlier — SDG-3"], 18)
pic(sl, "branding/journey.png", 8.5, 2.1, 3.9, 1.3)
cap(sl, 8.5, 3.6, 3.9, "Pilot: 500 families via 1 health centre. KPIs: scans, adherence %, alerts, referrals.", TEAL)

# 9 BUSINESS
sl = base(LIGHT, "Free for families forever — systems pay. Investors like that sentence.")
kick(sl, "9  BUSINESS MODEL & IMPLEMENTATION"); head(sl, "Free for families. Paid by systems.")
bullets(sl, 0.6, 1.95, 12.1, 4.6, ["Families: free 10 scans a month, Rs 99 a year for unlimited + Sehat File history", "Pharmacy kiosks at Rs 500 a month — scan, explain, print the timetable", "Health centres and blocks via NHM/CSR pilots — ASHA dashboards included", "Phase 2: Rs 299 beeper pill-box + BP-link hardware, funded by kiosk revenue"], 18)

# 10 FEASIBILITY
sl = base(LIGHT, "School team, real stack, honest risks with fixes. That honesty scores.")
kick(sl, "10  FEASIBILITY & SCALABILITY"); head(sl, "Built in weeks. Scales to Bharat.")
bullets(sl, 0.6, 1.95, 7.4, 4.9, ["App + verified DBs + tested AI engine — a school team built this in weeks, under Rs 2,000", "New language = one voice + one data column; Postgres scales to lakhs of families", "Messy writing fails → human confirm button; risky drugs → referral only, never a guess", "No internet → cached scans + offline voice; the expo-hall failure mode is designed out"], 18)
pic(sl, "branding/architecture.png", 8.5, 1.95, 3.9, 1.35)

# 11 ADVANTAGE
sl = base(LIGHT, "Read the table row by row. End on the dataset moat — 500 real parchis nobody else has.")
kick(sl, "11  COMPETITIVE ADVANTAGE"); head(sl, "A guardian, not another chatbot")
rows, cols = 4, 4
tsh = sl.shapes.add_table(rows, cols, Inches(0.6), Inches(1.95), Inches(12.1), Inches(3.3))
tbl = tsh.table
data = [["Feature", "Suraksha Parchi", "Health chatbot", "Scanner app"],
 ["Handwriting + mother-tongue voice", "YES, cited + pictograms", "NO", "Text only"],
 ["Duplicate + lab cross-check", "YES, with red flags", "NO", "NO"],
 ["Offline + ASHA workflow", "YES, cached + referral", "NO", "NO"]]
for r in range(rows):
    for c in range(cols):
        cell = tbl.cell(r, c); cell.text = data[r][c]
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(15); p.font.bold = (r == 0); p.font.color.rgb = WHITE if r == 0 else INK
        cell.fill.solid(); cell.fill.fore_color.rgb = TEAL if r == 0 else (GREY if r % 2 == 0 else WHITE)
cap(sl, 0.6, 5.5, 12.1, "Moat: 500 real Indian parchi scans (consented) + ASHA workflow + offline-first. Dataset + guardrailed prompts = IP.", TEAL)

# 12 TEAM
sl = base(INK, "Close warm: names, the ask (a pilot health centre), thank you in three languages.")
pic(sl, "branding/logo.png", 8.9, 1.6, 3.5, 3.5)
tx = sl.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(7.5), Inches(0.5))
p = tx.text_frame.paragraphs[0]; p.text = "12  TEAM & VISION"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = SAFFRON
tx2 = sl.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(7.5), Inches(1.0)); tx2.text_frame.word_wrap = True
p2 = tx2.text_frame.paragraphs[0]; p2.text = "CHANAKYA — Niti se Nirman tak"; p2.font.size = Pt(32); p2.font.bold = True; p2.font.color.rgb = WHITE
tx3 = sl.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(7.5), Inches(4.6)); tx3.text_frame.word_wrap = True; tf = tx3.text_frame
items = ["Ayush Pandey — Leader: product, app + AI engine, database design", "Kaushal Rawat — CMO: field testing, ASHA interviews, Bengali/Marathi content", "+2 teammates joining: designer + tester  •  Mentor: [To Fill]", "Ask: one health centre, 500 families, 4 weeks — then scale with you", "ayush4ru@gmail.com  •  github.com/Sentinal-Corporations/suraksha-sharthi"]
for i, it in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = "• " + it; p.font.size = Pt(17); p.font.color.rgb = WHITE; p.space_after = Pt(8)

out = f"{BASE}/ppt/CHANAKYA_SurakshaParchi.pptx"
prs.save(out); print("saved v3:", out, len(prs.slides), "slides")
